# -*- coding: utf-8 -*-
"""
core.inventory_db — Base de datos SQLite para Inventario, Control de Cajas/Lotes y Trazabilidad de ONUs.
Permite registrar cada ONU flasheada (PON, MAC, Modelo, Caja, Fechas, Duración, Resultado)
y generar reportes personalizados filtrados por día, modelo, caja o serial.
"""
import csv
import io
import json
import os
import re
import sqlite3
import time
from datetime import datetime, date

CORE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(CORE_DIR) if os.path.basename(CORE_DIR) == "core" else CORE_DIR
LOGS_DIR = os.path.join(PROJECT_ROOT, "logs")
CONFIG_DIR = os.path.join(PROJECT_ROOT, "config")
DB_PATH = os.path.join(LOGS_DIR, "inventario.db")
LEGACY_CSV = os.path.join(LOGS_DIR, "registro_flasheo.csv")

os.makedirs(LOGS_DIR, exist_ok=True)
os.makedirs(CONFIG_DIR, exist_ok=True)

_DB_INITIALIZED = False


def get_connection():
    """Retorna una conexión a la base de datos SQLite con row_factory configurado."""
    conn = sqlite3.connect(DB_PATH, timeout=20)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("PRAGMA journal_mode = WAL")
        conn.execute("PRAGMA synchronous = NORMAL")
    except Exception:
        pass
    return conn


def init_db():
    """Crea las tablas e índices si no existen."""
    global _DB_INITIALIZED
    if _DB_INITIALIZED:
        return

    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS lotes_cajas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                codigo_caja TEXT NOT NULL UNIQUE,
                modelo_id TEXT NOT NULL,
                modelo_nombre TEXT,
                cantidad_total INTEGER NOT NULL DEFAULT 20,
                cantidad_procesadas INTEGER DEFAULT 0,
                cantidad_exitosas INTEGER DEFAULT 0,
                cantidad_fallidas INTEGER DEFAULT 0,
                fecha_inicio TEXT NOT NULL,
                fecha_cierre TEXT,
                estado TEXT NOT NULL DEFAULT 'EN_PROCESO',
                observaciones TEXT DEFAULT ''
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS onus_flasheadas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                caja_id INTEGER,
                codigo_caja TEXT,
                modelo_id TEXT NOT NULL,
                pon_sn TEXT,
                pon_original TEXT,
                mac TEXT,
                puerto TEXT,
                ip TEXT,
                resultado TEXT NOT NULL,
                vlan_ok TEXT DEFAULT 'NO',
                firmware TEXT,
                detalles TEXT,
                fecha_hora TEXT NOT NULL,
                fecha_dia TEXT NOT NULL,
                duracion_segundos INTEGER DEFAULT 0,
                captura_path TEXT,
                FOREIGN KEY (caja_id) REFERENCES lotes_cajas(id) ON DELETE SET NULL
            )
        """)

        cursor.execute("CREATE INDEX IF NOT EXISTS idx_onus_fecha_dia ON onus_flasheadas(fecha_dia)")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_onus_caja ON onus_flasheadas(codigo_caja)")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_onus_modelo ON onus_flasheadas(modelo_id)")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_onus_pon ON onus_flasheadas(pon_sn)")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_onus_mac ON onus_flasheadas(mac)")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_onus_resultado ON onus_flasheadas(resultado)")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_cajas_estado ON lotes_cajas(estado)")

        conn.commit()

    _DB_INITIALIZED = True
    _check_and_migrate_legacy()


def format_vsol_pon(pon, mac=""):
    """Formatea el PON Serial al estándar GPON (VSOL00XXXXXX)."""
    if pon:
        p = str(pon).strip()
        if "-" in p:
            tail = p.split("-")[-1]
            if len(tail) >= 6:
                return "VSOL00" + tail[-6:].upper()
        elif len(p) >= 6:
            return "VSOL00" + p[-6:].upper()
    if mac:
        mc = re.sub(r"[^0-9A-Fa-f]", "", str(mac))
        if len(mc) >= 6:
            return "VSOL00" + mc[-6:].upper()
    return ""


def crear_o_abrir_caja(codigo_caja, modelo_id, cantidad_total=20, modelo_nombre=None, observaciones=""):
    """Crea una nueva caja o abre una existente si ya fue creada."""
    init_db()
    codigo_caja = str(codigo_caja).strip().upper()
    if not codigo_caja:
        codigo_caja = f"CJ-{datetime.now().strftime('%Y%m%d-%H%M%S')}"

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM lotes_cajas WHERE codigo_caja = ?", (codigo_caja,))
        existente = cursor.fetchone()

        if existente:
            caja_id = existente["id"]
            cursor.execute("""
                UPDATE lotes_cajas
                SET modelo_id = ?, modelo_nombre = coalesce(?, modelo_nombre),
                    cantidad_total = coalesce(?, cantidad_total),
                    estado = 'EN_PROCESO'
                WHERE id = ?
            """, (modelo_id, modelo_nombre, int(cantidad_total), caja_id))
        else:
            cursor.execute("""
                INSERT INTO lotes_cajas (codigo_caja, modelo_id, modelo_nombre, cantidad_total, fecha_inicio, estado, observaciones)
                VALUES (?, ?, ?, ?, ?, 'EN_PROCESO', ?)
            """, (codigo_caja, modelo_id, modelo_nombre or modelo_id, int(cantidad_total), now_str, observaciones))
            caja_id = cursor.lastrowid

        conn.commit()

    return obtener_caja_por_id(caja_id)


def obtener_caja_activa():
    """Retorna la caja que actualmente está marcada como 'EN_PROCESO' más reciente."""
    init_db()
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT * FROM lotes_cajas
            WHERE estado = 'EN_PROCESO'
            ORDER BY id DESC LIMIT 1
        """)
        row = cursor.fetchone()
        if row:
            return dict(row)
    return None


def obtener_caja_por_id(caja_id):
    init_db()
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM lotes_cajas WHERE id = ?", (caja_id,))
        row = cursor.fetchone()
        if row:
            return dict(row)
    return None


def obtener_caja_por_codigo(codigo_caja):
    init_db()
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM lotes_cajas WHERE codigo_caja = ?", (str(codigo_caja).strip().upper(),))
        row = cursor.fetchone()
        if row:
            return dict(row)
    return None


def cerrar_caja(caja_id_o_codigo, estado="COMPLETADA"):
    """Cierra una caja marcando su fecha de finalización."""
    init_db()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_connection() as conn:
        cursor = conn.cursor()
        if str(caja_id_o_codigo).isdigit():
            cursor.execute("""
                UPDATE lotes_cajas
                SET estado = ?, fecha_cierre = ?
                WHERE id = ?
            """, (estado, now_str, int(caja_id_o_codigo)))
        else:
            cursor.execute("""
                UPDATE lotes_cajas
                SET estado = ?, fecha_cierre = ?
                WHERE codigo_caja = ?
            """, (estado, now_str, str(caja_id_o_codigo).strip().upper()))
        conn.commit()
    return True


def cancelar_caja(caja_id_o_codigo, eliminar_si_vacia=True):
    """Cancela una caja activa. Si no tiene ONUs procesadas, la elimina por completo para no ensuciar el historial."""
    init_db()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_connection() as conn:
        cursor = conn.cursor()
        is_num = str(caja_id_o_codigo).isdigit()
        if is_num:
            cursor.execute("SELECT id, codigo_caja FROM lotes_cajas WHERE id = ?", (int(caja_id_o_codigo),))
        else:
            cursor.execute("SELECT id, codigo_caja FROM lotes_cajas WHERE codigo_caja = ?", (str(caja_id_o_codigo).strip().upper(),))
        row = cursor.fetchone()
        if not row:
            return False
        c_id = row["id"]
        c_cod = row["codigo_caja"]

        cursor.execute("SELECT COUNT(*) as count FROM onus_flasheadas WHERE caja_id = ? OR codigo_caja = ?", (c_id, c_cod))
        count = cursor.fetchone()["count"]

        if count == 0 and eliminar_si_vacia:
            cursor.execute("DELETE FROM lotes_cajas WHERE id = ?", (c_id,))
        else:
            cursor.execute("UPDATE lotes_cajas SET estado = 'CANCELADA', fecha_cierre = ? WHERE id = ?", (now_str, c_id))
        conn.commit()
    return True



def listar_cajas(limit=100):
    """Lista todas las cajas registradas ordenadas cronológicamente inverso."""
    init_db()
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT c.*,
                   (SELECT COUNT(*) FROM onus_flasheadas WHERE caja_id = c.id) as total_onus_db,
                   (SELECT COUNT(*) FROM onus_flasheadas WHERE caja_id = c.id AND resultado IN ('EXITO', 'YA_CONFIGURADA')) as total_ok_db
            FROM lotes_cajas c
            ORDER BY c.id DESC
            LIMIT ?
        """, (limit,))
        return [dict(r) for r in cursor.fetchall()]


def registrar_onu(modelo_id, pon, mac, puerto, ip, resultado, vlan_ok, detalles,
                  codigo_caja=None, firmware=None, duracion_segundos=0, captura_path=None,
                  fecha_hora=None):
    """Registra una ONU procesada en la base de datos y actualiza las métricas de la caja."""
    init_db()
    now = datetime.now()
    ts_str = fecha_hora or now.strftime("%Y-%m-%d %H:%M:%S")
    dia_str = ts_str.split(" ")[0] if " " in ts_str else now.strftime("%Y-%m-%d")

    pon_orig = str(pon).strip() if pon else ""
    mac_val = str(mac).strip().upper() if mac else ""
    pon_form = format_vsol_pon(pon_orig, mac_val)

    caja = None
    if codigo_caja:
        caja = obtener_caja_por_codigo(codigo_caja)
    if not caja:
        caja = obtener_caja_activa()

    caja_id = caja["id"] if caja else None
    cod_caja_final = caja["codigo_caja"] if caja else (codigo_caja or "SIN_CAJA")

    is_sin_onu = resultado == "SIN_ONU"

    with get_connection() as conn:
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO onus_flasheadas (
                caja_id, codigo_caja, modelo_id, pon_sn, pon_original, mac,
                puerto, ip, resultado, vlan_ok, firmware, detalles,
                fecha_hora, fecha_dia, duracion_segundos, captura_path
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            caja_id, cod_caja_final, modelo_id, pon_form, pon_orig, mac_val,
            str(puerto), str(ip), resultado, vlan_ok, firmware or "", str(detalles),
            ts_str, dia_str, int(duracion_segundos or 0), captura_path or ""
        ))
        onu_id = cursor.lastrowid

        if caja_id and not is_sin_onu:
            cursor.execute("""
                UPDATE lotes_cajas
                SET cantidad_procesadas = (SELECT COUNT(*) FROM onus_flasheadas WHERE caja_id = ? AND resultado != 'SIN_ONU'),
                    cantidad_exitosas = (SELECT COUNT(*) FROM onus_flasheadas WHERE caja_id = ? AND resultado IN ('EXITO', 'YA_CONFIGURADA', 'CHECK_OK')),
                    cantidad_fallidas = (SELECT COUNT(*) FROM onus_flasheadas WHERE caja_id = ? AND resultado IN ('ERROR', 'PARCIAL', 'FAIL'))
                WHERE id = ?
            """, (caja_id, caja_id, caja_id, caja_id))

            cursor.execute("SELECT cantidad_total, cantidad_procesadas FROM lotes_cajas WHERE id = ?", (caja_id,))
            c_row = cursor.fetchone()
            if c_row and c_row["cantidad_procesadas"] >= c_row["cantidad_total"]:
                cursor.execute("UPDATE lotes_cajas SET estado = 'COMPLETADA' WHERE id = ? AND estado = 'EN_PROCESO'", (caja_id,))

        conn.commit()

    return onu_id


def consultar_onus_por_caja(codigo_caja_o_id):
    """Consulta todas las ONUs procesadas de una caja específica ordenadas por llegada."""
    init_db()
    with get_connection() as conn:
        cursor = conn.cursor()
        if str(codigo_caja_o_id).isdigit():
            cursor.execute("""
                SELECT o.*, c.modelo_nombre, c.cantidad_total as caja_cantidad_total
                FROM onus_flasheadas o
                LEFT JOIN lotes_cajas c ON o.caja_id = c.id
                WHERE o.caja_id = ? AND o.resultado != 'SIN_ONU'
                ORDER BY o.id ASC
            """, (int(codigo_caja_o_id),))
        else:
            cursor.execute("""
                SELECT o.*, c.modelo_nombre, c.cantidad_total as caja_cantidad_total
                FROM onus_flasheadas o
                LEFT JOIN lotes_cajas c ON o.caja_id = c.id
                WHERE o.codigo_caja = ? AND o.resultado != 'SIN_ONU'
                ORDER BY o.id ASC
            """, (str(codigo_caja_o_id).strip().upper(),))

        return [dict(r) for r in cursor.fetchall()]


def consultar_reporte(fecha_desde=None, fecha_hasta=None, modelo=None, codigo_caja=None,
                      resultado=None, search=None, limit=1000, offset=0, excluir_sin_onu=True):
    """Consulta flexible con filtros múltiples para la UI y exportaciones."""
    init_db()
    query = """
        SELECT o.*, c.modelo_nombre, c.cantidad_total as caja_cantidad_total
        FROM onus_flasheadas o
        LEFT JOIN lotes_cajas c ON o.caja_id = c.id
        WHERE 1=1
    """
    params = []

    if excluir_sin_onu:
        query += " AND o.resultado != 'SIN_ONU'"

    if fecha_desde:
        query += " AND o.fecha_dia >= ?"
        params.append(str(fecha_desde).strip())

    if fecha_hasta:
        query += " AND o.fecha_dia <= ?"
        params.append(str(fecha_hasta).strip())

    if modelo and modelo != "TODOS":
        query += " AND o.modelo_id = ?"
        params.append(str(modelo).strip())

    if codigo_caja and codigo_caja != "TODAS":
        query += " AND o.codigo_caja = ?"
        params.append(str(codigo_caja).strip().upper())

    if resultado and resultado != "TODOS":
        if resultado == "EXITO":
            query += " AND o.resultado IN ('EXITO', 'YA_CONFIGURADA')"
        elif resultado == "FALLO":
            query += " AND o.resultado IN ('ERROR', 'PARCIAL', 'FAIL')"
        else:
            query += " AND o.resultado = ?"
            params.append(str(resultado).strip())

    if search:
        s = f"%{str(search).strip()}%"
        query += " AND (o.pon_sn LIKE ? OR o.mac LIKE ? OR o.codigo_caja LIKE ? OR o.detalles LIKE ? OR o.ip LIKE ?)"
        params.extend([s, s, s, s, s])

    query += " ORDER BY o.id DESC LIMIT ? OFFSET ?"
    params.extend([int(limit), int(offset)])

    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(query, params)
        rows = cursor.fetchall()
        return [dict(r) for r in rows]


def obtener_estadisticas_resumen(fecha_dia=None):
    """Retorna métricas generales para los contadores y KPIs de la UI."""
    init_db()
    dia = fecha_dia or datetime.now().strftime("%Y-%m-%d")
    with get_connection() as conn:
        cursor = conn.cursor()

        cursor.execute("""
            SELECT
                COUNT(*) as total_hoy,
                SUM(CASE WHEN resultado IN ('EXITO', 'YA_CONFIGURADA') THEN 1 ELSE 0 END) as ok_hoy,
                SUM(CASE WHEN resultado IN ('ERROR', 'PARCIAL', 'FAIL') THEN 1 ELSE 0 END) as fallos_hoy
            FROM onus_flasheadas
            WHERE fecha_dia = ? AND resultado != 'SIN_ONU'
        """, (dia,))
        hoy = cursor.fetchone()

        cursor.execute("""
            SELECT
                COUNT(*) as total_historico,
                SUM(CASE WHEN resultado IN ('EXITO', 'YA_CONFIGURADA') THEN 1 ELSE 0 END) as ok_historico,
                SUM(CASE WHEN resultado IN ('ERROR', 'PARCIAL', 'FAIL') THEN 1 ELSE 0 END) as fallos_historico
            FROM onus_flasheadas
            WHERE resultado != 'SIN_ONU'
        """)
        historico = cursor.fetchone()

        cursor.execute("SELECT COUNT(*) as total_cajas FROM lotes_cajas")
        cajas = cursor.fetchone()

        return {
            "hoy": {
                "total": hoy["total_hoy"] or 0,
                "exitosas": hoy["ok_hoy"] or 0,
                "fallidas": hoy["fallos_hoy"] or 0,
                "tasa_exito": round(((hoy["ok_hoy"] or 0) / hoy["total_hoy"] * 100), 1) if hoy["total_hoy"] else 0.0
            },
            "historico": {
                "total": historico["total_historico"] or 0,
                "exitosas": historico["ok_historico"] or 0,
                "fallidas": historico["fallos_historico"] or 0,
                "tasa_exito": round(((historico["ok_historico"] or 0) / historico["total_historico"] * 100), 1) if historico["total_historico"] else 0.0,
                "total_cajas": cajas["total_cajas"] or 0
            }
        }


def exportar_csv(filtros=None):
    """Genera un stream CSV estructurado a partir de los filtros seleccionados."""
    records = consultar_reporte(**(filtros or {}), limit=10000)
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")

    writer.writerow([
        "FechaHora", "CodigoCaja", "Modelo", "PON_GPON_SN", "PON_Original",
        "MAC_Address", "Puerto", "IP", "Resultado", "VLAN3_OK", "Firmware",
        "Duracion_Segundos", "Detalles"
    ])

    for r in records:
        writer.writerow([
            r.get("fecha_hora", ""),
            r.get("codigo_caja", ""),
            r.get("modelo_id", ""),
            r.get("pon_sn", ""),
            r.get("pon_original", ""),
            r.get("mac", ""),
            r.get("puerto", ""),
            r.get("ip", ""),
            r.get("resultado", ""),
            r.get("vlan_ok", ""),
            r.get("firmware", ""),
            r.get("duracion_segundos", 0),
            r.get("detalles", "")
        ])

    return output.getvalue().encode("utf-8-sig")


def exportar_excel(filtros=None):
    """Genera un archivo Excel (.xlsx) con formato profesional listo para el sistema de inventario."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return exportar_csv(filtros), "text/csv", "reporte_onus.csv"

    records = consultar_reporte(**(filtros or {}), limit=10000)
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario ONUs"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    headers = [
        "Fecha / Hora", "Caja / Lote", "Modelo ONU", "PON Serial (GPON SN)",
        "MAC Address", "Puerto", "IP Gestión", "Resultado", "VLAN 3",
        "Firmware", "Tiempo (s)", "Detalles"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    ws.row_dimensions[1].height = 28

    ok_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_data = Font(name="Calibri", size=10)

    for row_idx, r in enumerate(records, start=2):
        res = r.get("resultado", "")
        is_ok = res in ("EXITO", "YA_CONFIGURADA", "CHECK_OK")

        row_data = [
            r.get("fecha_hora", ""),
            r.get("codigo_caja", ""),
            r.get("modelo_id", ""),
            r.get("pon_sn", ""),
            r.get("mac", ""),
            r.get("puerto", ""),
            r.get("ip", ""),
            res,
            r.get("vlan_ok", ""),
            r.get("firmware", ""),
            r.get("duracion_segundos", 0),
            r.get("detalles", "")
        ]
        ws.append(row_data)

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_data
            cell.border = thin_border
            if col_idx in [1, 2, 4, 5, 6, 7, 8, 9, 11]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            if col_idx == 8:
                cell.fill = ok_fill if is_ok else fail_fill

        ws.row_dimensions[row_idx].height = 20

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"Inventario_ONUs_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"


def _check_and_migrate_legacy():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) as c FROM onus_flasheadas")
        if cursor.fetchone()["c"] > 0:
            return

    legacy_candidates = [
        LEGACY_CSV,
        os.path.join(PROJECT_ROOT, "registro_flasheo.csv")
    ]
    csv_file = None
    for cand in legacy_candidates:
        if os.path.exists(cand):
            csv_file = cand
            break

    if not csv_file:
        return

    try:
        with open(csv_file, "r", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.DictReader(f)
            caja_default = "LOTE-HISTORICO"
            now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            with get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute("""
                    INSERT OR IGNORE INTO lotes_cajas (codigo_caja, modelo_id, modelo_nombre, cantidad_total, fecha_inicio, estado, observaciones)
                    VALUES (?, 'V2801S-B', 'Lote Histórico Migrado', 9999, ?, 'COMPLETADA', 'Migración automática desde CSV')
                """, (caja_default, now_str))

                cursor.execute("SELECT id FROM lotes_cajas WHERE codigo_caja = ?", (caja_default,))
                c_row = cursor.fetchone()
                caja_id = c_row["id"] if c_row else None

                for row in reader:
                    ts = row.get("FechaHora") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    dia = ts.split(" ")[0] if " " in ts else datetime.now().strftime("%Y-%m-%d")
                    p_orig = row.get("PON/Serial") or ""
                    mac_v = row.get("MAC") or ""
                    p_form = row.get("PON Formateado") or format_vsol_pon(p_orig, mac_v)
                    res = row.get("Resultado") or "DESCONOCIDO"

                    cursor.execute("""
                        INSERT INTO onus_flasheadas (
                            caja_id, codigo_caja, modelo_id, pon_sn, pon_original, mac,
                            puerto, ip, resultado, vlan_ok, detalles, fecha_hora, fecha_dia
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        caja_id, caja_default, "V2801S-B", p_form, p_orig, mac_v,
                        row.get("Puerto/ID") or "1", row.get("IP") or "",
                        res, row.get("VLAN3 Verificada") or "NO",
                        row.get("Detalles") or "", ts, dia
                    ))
                conn.commit()
    except Exception as ex:
        print(f"[InventoryDB] Error al migrar histórico legacy: {ex}")


if __name__ == "__main__":
    init_db()
    print("Base de datos de inventario inicializada correctamente en:", DB_PATH)
