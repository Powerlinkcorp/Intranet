import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    cell.width = Inches(6.5)
    
    # Left border accent
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="3B82F6"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(31, 41, 55)

def add_callout(doc, text, title="NOTA IMPORTANTE", border_color="EAB308", bg_color="FEF9C3"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    cell.width = Inches(6.5)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"⚠️ {title}: ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(133, 77, 14)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(113, 63, 18)

def create_document():
    doc = docx.Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Styles config
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 51, 51)

    # Document Header / Banner
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("INSTRUCTIVO OPERATIVO OFICIAL")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(26, 26, 26)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    sub_run = sub_p.add_run("Flujo de Trabajo Git, Integración a Pruebas y Despliegues en Servidores (QA y Producción)")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(202, 138, 4) # Dark Gold Powerlink
    sub_run.bold = True

    # Metadata Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Proyecto:", " Intranet Powerlink"), ("Fecha:", " Octubre 2026")],
        [("Servidor Destino:", " 10.0.1.179 (Ubuntu 24.04)"), ("Autorizado por:", " Administrador de Integración")]
    ]
    for row_idx, row in enumerate(meta_data):
        for col_idx, (k, v) in enumerate(row):
            cell = meta_table.cell(row_idx, col_idx)
            set_cell_background(cell, "F9FAFB")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r1 = p.add_run(k)
            r1.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = RGBColor(75, 85, 99)
            r2 = p.add_run(v)
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = RGBColor(17, 24, 39)
    
    doc.add_paragraph() # Spacer

    # 1. INTRODUCCIÓN
    h1 = doc.add_heading("1. Propósito y Arquitectura de Ramas", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(17, 24, 39)
        r.bold = True

    p = doc.add_paragraph(
        "Este instructivo define la metodología estándar obligatoria para todo el equipo de desarrollo de la Intranet. "
        "El objetivo principal es eliminar riesgos de indisponibilidad del servidor causados por conflictos no resueltos, "
        "archivos residuales o subidas directas no controladas."
    )
    p.paragraph_format.space_after = Pt(8)

    # Branches Table
    table = doc.add_table(rows=5, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Rama Git", "Propósito", "Responsable", "Entorno Asociado"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_background(cell, "1E293B")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)

    branches_data = [
        ("<mi-rama> (ej: whilly, dixon)", "Desarrollo local activo de tareas asignadas", "Desarrollador", "Entorno Local"),
        ("pruebas", "Rama común de integración de código del equipo", "Todos los desarrolladores", "Validación integrada local"),
        ("qa", "Aseguramiento de Calidad y pruebas en servidor", "Solo Administrador", "Servidor QA (10.0.1.179:8080)"),
        ("Produccion / main", "Entorno productivo para toda la organización", "Solo Administrador", "Servidor Prod (10.0.1.179:80)")
    ]

    for row_idx, data in enumerate(branches_data, start=1):
        bg = "FFFFFF" if row_idx % 2 == 1 else "F9FAFB"
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            if col_idx == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(30, 64, 175)

    doc.add_paragraph() # Spacer

    add_callout(
        doc,
        "Ningún desarrollador debe realizar 'push' directo a las ramas qa, Produccion o main. "
        "Todo cambio debe originarse en la rama personal, ser integrado y probado en 'pruebas', y únicamente el Administrador "
        "está facultado para desplegar hacia los entornos de QA y Producción.",
        title="REGLA FUNDAMENTAL DE SEGURIDAD",
        border_color="DC2626",
        bg_color="FEF2F2"
    )

    # 2. GUÍA PARA EL DESARROLLADOR
    h2 = doc.add_heading("2. Guía para el Desarrollador (Ciclo de Trabajo)", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(17, 24, 39)
        r.bold = True

    # 2.1
    h2_1 = doc.add_heading("2.1 Preparación y Lista Negra antes del Commit", level=2)
    h2_1.paragraph_format.space_before = Pt(10)
    h2_1.paragraph_format.space_after = Pt(4)
    
    doc.add_paragraph(
        "Al finalizar una funcionalidad o corrección en tu rama local, ejecuta un análisis de estado. "
        "Bajo ninguna circunstancia incluyas archivos de entorno o bases de datos locales en tu commit:"
    )
    
    bullet_items = [
        ("Archivos de Configuración y Claves: ", ".env (contiene credenciales maestras y nunca debe subirse)."),
        ("Bases de Datos Locales / Temporales: ", "intranet_local.db, *.sqlite3, *.db-shm, *.db-wal."),
        ("Caché de Python: ", "__pycache__/, *.pyc."),
        ("Entornos Virtuales y Registros: ", "venv/, *.log, directorios temporales scratch/.")
    ]
    for bold_txt, norm_txt in bullet_items:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after = Pt(2)
        r_bold = bp.add_run(bold_txt)
        r_bold.bold = True
        r_bold.font.color.rgb = RGBColor(185, 28, 28)
        bp.add_run(norm_txt)

    doc.add_paragraph("Valida que tu código compile sin errores de sintaxis antes de preparar el commit:")
    add_code_block(doc, "python -m py_compile main.py routers/*.py")

    # 2.2
    h2_2 = doc.add_heading("2.2 Creación del Commit Local Selectivo", level=2)
    h2_2.paragraph_format.space_before = Pt(10)
    h2_2.paragraph_format.space_after = Pt(4)
    doc.add_paragraph(
        "Agrega exclusivamente los archivos involucrados en tu tarea específica. "
        "Evita el comando 'git add .' si existen archivos temporales o bases de datos modificadas en tu directorio:"
    )
    add_code_block(doc, "git add routers/operaciones_router.py templates/operaciones.html\ngit commit -m \"feat(operaciones): agregar filtro avanzado de validacion ozmap\"")

    # 2.3
    h2_3 = doc.add_heading("2.3 Sincronización y Subida a tu Rama Remota", level=2)
    h2_3.paragraph_format.space_before = Pt(10)
    h2_3.paragraph_format.space_after = Pt(4)
    doc.add_paragraph("Asegura tu trabajo en GitHub subiendo a tu propia rama personal:")
    add_code_block(doc, "# 1. Traer posibles cambios remotos previos de tu rama\ngit pull origin <tu-rama>\n\n# 2. Enviar tu trabajo a tu rama remota\ngit push origin <tu-rama>")

    # 2.4
    h2_4 = doc.add_heading("2.4 Integración hacia la Rama 'pruebas'", level=2)
    h2_4.paragraph_format.space_before = Pt(10)
    h2_4.paragraph_format.space_after = Pt(4)
    doc.add_paragraph(
        "Una vez que tus cambios están respaldados en tu rama remota, incorpóralos en la rama de integración 'pruebas':"
    )
    add_code_block(doc, "# 1. Obtener ramas remotas actualizadas\ngit fetch origin\n\n# 2. Cambiar a pruebas y actualizarla\ngit checkout pruebas\ngit pull origin pruebas\n\n# 3. Fusionar tu rama en pruebas\ngit merge <tu-rama>\n\n# 4. Verificar ausencia de marcadores y compilar\ngit diff --check\npython -m py_compile main.py routers/*.py\n\n# 5. Enviar a GitHub\ngit push origin pruebas\n\n# 6. Regresar a tu rama de trabajo\ngit checkout <tu-rama>")

    add_callout(
        doc,
        "Si durante el merge Git reporta conflictos, resuélvelos minuciosamente en el editor. "
        "NUNCA confirmes un commit que contenga etiquetas residuales (<<<<<<< HEAD, =======, >>>>>>>). "
        "Esos marcadores generan fallos críticos 'SyntaxError: invalid decimal literal' al arrancar Uvicorn.",
        title="OBLIGACIÓN EN RESOLUCIÓN DE CONFLICTOS",
        border_color="F59E0B",
        bg_color="FFFBEB"
    )

    # 3. GUÍA PARA EL ADMINISTRADOR
    h3 = doc.add_heading("3. Guía Exclusiva del Administrador (Despliegues en Servidores)", level=1)
    h3.paragraph_format.space_before = Pt(18)
    h3.paragraph_format.space_after = Pt(6)
    for r in h3.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(17, 24, 39)
        r.bold = True

    p = doc.add_paragraph(
        "El Administrador es la única persona autorizada a manipular las ramas 'qa' y 'Produccion', "
        "así como a reiniciar los servicios en el servidor Linux 10.0.1.179."
    )

    # 3.1
    h3_1 = doc.add_heading("3.1 Promover de 'pruebas' hacia 'qa'", level=2)
    h3_1.paragraph_format.space_before = Pt(10)
    h3_1.paragraph_format.space_after = Pt(4)
    doc.add_paragraph("En la máquina del Administrador, consolidar los cambios validados en 'pruebas' dentro de 'qa':")
    add_code_block(doc, "git checkout qa\ngit pull origin qa\ngit merge pruebas\npython -m py_compile main.py routers/*.py\ngit push origin qa")

    # 3.2
    h3_2 = doc.add_heading("3.2 Despliegue en el Servidor de QA (Puerto 8080)", level=2)
    h3_2.paragraph_format.space_before = Pt(10)
    h3_2.paragraph_format.space_after = Pt(4)
    doc.add_paragraph("Conectarse al servidor por SSH y aplicar la actualización en el entorno QA:")
    add_code_block(doc, "# 1. Conexión SSH\nssh root@10.0.1.179\n\n# 2. Descargar cambios en QA\ncd /var/www/intranet_qa\ngit status\ngit pull origin qa\n\n# 3. Verificar que no existan conflictos ni errores\ngrep -rn --exclude-dir={.git,venv,node_modules} '<<<<<<<' /var/www/intranet_qa\n/var/www/intranet_qa/venv/bin/python -m py_compile main.py routers/*.py\n\n# 4. Reiniciar servicio y auditar\nsystemctl restart intranet_qa\nsystemctl status intranet_qa --no-pager\ncurl -I http://127.0.0.1:8001/login\n\n# 5. Probar acceso web en: http://10.0.1.179:8080")

    # 3.3
    h3_3 = doc.add_heading("3.3 Despliegue en el Servidor de Producción (Puerto 80)", level=2)
    h3_3.paragraph_format.space_before = Pt(10)
    h3_3.paragraph_format.space_after = Pt(4)
    doc.add_paragraph(
        "Una vez que el equipo certifica el funcionamiento en QA, se promueve a Producción "
        "siguiendo la regla estricta de respaldo previo:"
    )
    add_code_block(doc, "# PASO A: Respaldo Preventivo Obligatorio (Rollback inmediato)\nBACKUP_PROD=/var/www/intranet_prod_backup_$(date +%Y%m%d_%H%M%S)\ncp -a /var/www/intranet_prod $BACKUP_PROD\necho \"Respaldo creado en: $BACKUP_PROD\"\n\n# PASO B: Sincronización Segura QA -> Producción\nrsync -av --delete \\\n  --exclude='venv' \\\n  --exclude='.env' \\\n  --exclude='logs' \\\n  --exclude='static/uploads' \\\n  /var/www/intranet_qa/ /var/www/intranet_prod/\n\n# PASO C: Garantizar Permisos de Apache\nchown -R root:www-data /var/www/intranet_prod\n\n# PASO D: Validar compilación en Producción\ncd /var/www/intranet_prod\n/var/www/intranet_prod/venv/bin/python -m py_compile main.py routers/*.py\n\n# PASO E: Reiniciar servicio y verificar respuesta\nsystemctl restart intranet_prod\nsystemctl status intranet_prod --no-pager\ncurl -I http://127.0.0.1:8000/login\ncurl -I http://10.0.1.179:80/login")

    # 4. PROTOCOLO DE EMERGENCIA
    h4 = doc.add_heading("4. Protocolo de Diagnóstico y Recuperación de Emergencia", level=1)
    h4.paragraph_format.space_before = Pt(18)
    h4.paragraph_format.space_after = Pt(6)
    for r in h4.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(17, 24, 39)
        r.bold = True

    doc.add_paragraph(
        "Si tras una actualización el servidor responde con '502 Bad Gateway' o '503 Service Unavailable', "
        "ejecuta la siguiente secuencia de resolución inmediata:"
    )
    add_code_block(doc, "# 1. Identificar el servicio afectado\nsystemctl status intranet_qa --no-pager\nsystemctl status intranet_prod --no-pager\n\n# 2. Inspeccionar la causa del fallo en los registros de Uvicorn\njournalctl -u intranet_qa --no-pager -n 50\n# o para producción:\njournalctl -u intranet_prod --no-pager -n 50\n\n# 3. Rastrear marcadores de conflicto de Git\ngrep -rn --exclude-dir={.git,venv,node_modules} '<<<<<<<' /var/www/intranet_qa\n\n# 4. Rollback de Emergencia en Producción (si el fallo persiste)\nsystemctl stop intranet_prod\nrsync -av --delete /var/www/<directorio_backup>/ /var/www/intranet_prod/\nsystemctl restart intranet_prod")

    # 5. HABILIDAD ANTIGRAVITY
    h5 = doc.add_heading("5. Integración con Antigravity: Habilidad 'git'", level=1)
    h5.paragraph_format.space_before = Pt(18)
    h5.paragraph_format.space_after = Pt(6)
    for r in h5.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(17, 24, 39)
        r.bold = True

    doc.add_paragraph(
        "Para simplificar y automatizar este flujo, se ha configurado la habilidad 'git' en el entorno de Antigravity "
        "(.agents/skills/git/SKILL.md). Puedes solicitar su ejecución directa en el chat mediante comandos naturales:"
    )

    skill_cmds = [
        ("Subir cambios propios: ", "'Antigravity, usa la habilidad git para subir mis cambios actuales a mi rama y luego a pruebas'"),
        ("Despliegue a QA (Admin): ", "'Antigravity, como administrador actualiza qa desde pruebas y despliega en el servidor'"),
        ("Despliegue a Producción (Admin): ", "'Antigravity, clona y despliega QA a Producción con respaldo preventivo'"),
        ("Auditoría de Salud: ", "'Antigravity, verifica el estado y respuesta de los servicios de QA y Producción'")
    ]
    for bold_txt, norm_txt in skill_cmds:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after = Pt(2)
        r_bold = bp.add_run(bold_txt)
        r_bold.bold = True
        r_bold.font.color.rgb = RGBColor(30, 64, 175)
        bp.add_run(norm_txt)

    # Save
    out_path = "Instructivo_Git_y_Despliegues_Powerlink.docx"
    doc.save(out_path)
    print(f"Documento generado exitosamente en: {out_path}")

if __name__ == "__main__":
    create_document()
